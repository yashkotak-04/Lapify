import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for name, value in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{name}')
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_speech_document(filename):
    doc = Document()

    # Set page margins (0.8 inch)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Base colors
    c_primary = RGBColor(37, 99, 235)      # Royal Blue (#2563eb)
    c_dark = RGBColor(15, 23, 42)          # Slate Dark (#0f172a)
    c_slate = RGBColor(71, 85, 105)        # Slate Grey (#475569)
    c_white = RGBColor(255, 255, 255)

    # --- Title Banner ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("🏆 LAPIFY – TOP 20 EXHIBITION FAIR")
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = c_primary

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Master Presentation Speech, Live Demo Walkthrough & Jury Viva Guide")
    sub_run.font.name = 'Arial'
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = c_slate

    doc.add_paragraph() # Spacer

    # Helper function for Headings
    def add_section_heading(title_text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(title_text)
        r.font.name = 'Arial'
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = c_primary
        return h

    def add_subheading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = c_dark
        return p

    def add_body_p(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix + " ")
            r_pre.font.name = 'Arial'
            r_pre.font.size = Pt(10.5)
            r_pre.font.bold = True
            r_pre.font.color.rgb = c_dark
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(10.5)
        r.font.italic = italic
        r.font.color.rgb = c_dark
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix + " ")
            r_pre.font.name = 'Arial'
            r_pre.font.size = Pt(10.5)
            r_pre.font.bold = True
            r_pre.font.color.rgb = c_dark
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(10.5)
        r.font.color.rgb = c_dark
        return p

    # ==========================================
    # SECTION 1: THE 3-MINUTE MASTER SPEECH
    # ==========================================
    add_section_heading("🎙️ SECTION 1: THE 3-MINUTE MASTER PRESENTATION SPEECH")
    add_body_p("(Deliver this speech with confidence, clear voice modulation, and eye contact with evaluators)", italic=True)

    add_body_p('"Respected Teachers, Evaluators, and Guests,', bold_prefix="Opening Greeting:")
    add_body_p('Good morning/afternoon. My name is [Your Name], and today I am honored to present Lapify — a full-stack, secure marketplace and management ecosystem engineered specifically for buying, selling, and fulfilling certified new and pre-owned laptops.')

    add_subheading("🚩 The Real-World Problem")
    add_body_p('Today, purchasing or selling laptops online is filled with critical issues:')
    add_bullet('General e-commerce giants lack precision spec-filtering (e.g. searching specifically by processor generation, RAM expandability, or certified condition grading).', bold_prefix='1. Lack of Spec Focus:')
    add_bullet('Online classifieds (like OLX) suffer from rampant scams, absence of quality inspection, lack of buyer protection, and no formal tax invoicing.', bold_prefix='2. Zero Trust & Security:')
    add_bullet('Independent laptop sellers face complex listing hurdles and inventory management friction.', bold_prefix='3. High Seller Friction:')

    add_subheading("💡 The Solution: Lapify")
    add_body_p('Lapify solves these challenges through a unified three-pillar architecture:')

    add_bullet('Buyers enjoy instant faceted filtering (by brand, CPU, RAM, storage, condition, and price slider), a streamlined 4-step checkout with live quantity adjustments, real-time lifecycle order tracking, and instant automated PDF tax invoices with smart print fallback.', bold_prefix='1. For Buyers:')
    add_bullet('Sellers can publish listings with automatic technical description generation, condition grading, photo uploads, multi-unit stock controls, and transparent real-time approval status tracking with admin feedback notes.', bold_prefix='2. For Sellers:')
    add_bullet('Administrators manage a robust quality moderation desk to approve/reject listings before going live, manage fulfillment phases, resolve customer support inquiries, and analyze monthly platform revenue.', bold_prefix='3. For Administrators:')

    add_subheading("⚙️ Engineering & Security Highlights")
    add_bullet('Backend developed in PHP with MySQL relational schema, incorporating ACID-compliant transactions and pessimistic row locking (SELECT ... FOR UPDATE) to prevent race conditions or double-selling.', bold_prefix='• Database Concurrency:')
    add_bullet('100% Prepared Statements for total SQL injection immunity, CSRF token validation on all mutating requests, XSS HTML entity sanitization, and Bcrypt password hashing.', bold_prefix='• Multi-Layer Security:')
    add_bullet('Custom CSS3 design system with dark/light theme awareness, responsive flexbox/grid layouts, micro-animations, and fast page loads.', bold_prefix='• Responsive Design System:')

    add_body_p('In summary, Lapify is not just a demo concept — it is a production-ready, highly reliable commercial marketplace engineered to make laptop trade transparent, secure, and accessible to everyone. Thank you, and I look forward to demonstrating our live platform for you!"', bold_prefix="Conclusion:")

    doc.add_paragraph() # Spacer

    # ==========================================
    # SECTION 2: 60-SECOND ELEVATOR PITCH
    # ==========================================
    add_section_heading("⚡ SECTION 2: THE 60-SECOND QUICK ELEVATOR PITCH")
    add_body_p("(Use this for quick evaluator stops or casual visitor walk-ins)", italic=True)

    add_body_p('"Hi! Welcome to Lapify! Have you ever tried buying a refurbished laptop online and worried about fake specs, unfair pricing, or getting an official invoice? Lapify solves that completely.')
    add_body_p('Lapify is an end-to-end specialized laptop marketplace connecting verified sellers and buyers with an admin-moderated catalog. Key features include precision hardware filtering, an interactive 4-step checkout with stock concurrency locking, instant PDF tax invoice downloads, and live order tracking.')
    add_body_p('Let me give you a quick 10-second live demo of placing an order and generating an invoice!"')

    doc.add_paragraph() # Spacer

    # ==========================================
    # SECTION 3: LIVE DEMO SCRIPT
    # ==========================================
    add_section_heading("🖥️ SECTION 3: STEP-BY-STEP LIVE DEMONSTRATION SCRIPT")
    add_body_p("Follow this exact step-by-step sequence while operating the screen:", italic=True)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["Step & Page", "Action on Screen", "What to Say to the Judges"]
    widths = [Inches(1.5), Inches(2.2), Inches(3.1)]

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], '2563eb')
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = 'Arial'
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = c_white

    demo_steps = [
        ("1. Storefront\n(index.php & buy.php)", "Open catalog, interact with CPU, RAM, and price slider filters.", "Here is our precision storefront. Users can instantly narrow down laptops by processor, RAM, and condition with zero lag."),
        ("2. Laptop Details\n(laptop-details.php)", "Open any product card; highlight specs, seller details & stock badge.", "Each listing displays verified technical specifications, condition grading, seller contact details, and live available stock count."),
        ("3. Checkout Cart\n(checkout_cart.php)", "Click 'Buy Now'. Click '+' on quantity stepper up to stock limit.", "Our multi-step checkout dynamically enforces inventory limits. Users can adjust quantity up to published stock, apply promo codes like LAPIFY50, and view live subtotal recalculations."),
        ("4. Payment & Order\n(checkout_payment.php)", "Fill shipping info, choose Cash on Delivery, click 'Place Order'.", "We execute atomic database locking (SELECT ... FOR UPDATE) to safely decrement stock and generate an order reference with live tracking."),
        ("5. Invoice Engine\n(invoice.php)", "Click 'Download Invoice PDF'. Open downloaded PDF.", "The system immediately compiles and streams a clean, branded PDF tax invoice with full tax itemization and automatic print fallback."),
        ("6. Admin Center\n(admin/dashboard.php)", "Switch to Admin tab. Show pending approvals & order status tracker.", "Admins have full command: inspecting pending listings before publication, tracking monthly marketplace revenue, and updating order fulfillment phases.")
    ]

    for step_title, action_text, speech_text in demo_steps:
        row = table.add_row()
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.width = widths[0]
        c1.width = widths[1]
        c2.width = widths[2]

        for c in (c0, c1, c2):
            set_cell_margins(c, top=100, bottom=100, left=120, right=120)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(step_title)
        r0.font.name = 'Arial'
        r0.font.size = Pt(9.5)
        r0.font.bold = True

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(action_text)
        r1.font.name = 'Arial'
        r1.font.size = Pt(9.5)

        p2 = c2.paragraphs[0]
        r2 = p2.add_run(speech_text)
        r2.font.name = 'Arial'
        r2.font.size = Pt(9.5)
        r2.font.italic = True

    doc.add_paragraph() # Spacer

    # ==========================================
    # SECTION 4: VIVA & JURY QUESTIONS
    # ==========================================
    add_section_heading("🧠 SECTION 4: TOP EVALUATOR VIVA QUESTIONS & HIGH-IMPACT ANSWERS")

    qa_list = [
        ("Q1: How do you handle race conditions when multiple users attempt to buy the last laptop simultaneously?",
         "Answer: We implement ACID-compliant transactions in MySQL with pessimistic row-level locking using 'SELECT ... FOR UPDATE'. When a checkout is submitted, the specific laptop record is locked exclusively until the order record is committed and stock is decremented. Any concurrent transaction waits safely, preventing overselling or stock anomalies."),
        
        ("Q2: What security measures protect Lapify against malicious attacks?",
         "Answer: We enforce multi-layer security: 1) Parameterized Prepared Statements across 100% of queries for SQL injection immunity; 2) Cryptographic CSRF tokens on all mutating HTTP forms; 3) Full input/output sanitization with htmlspecialchars against Cross-Site Scripting (XSS); and 4) Industry-standard Bcrypt password hashing."),
        
        ("Q3: How does the PDF invoice engine avoid 500 server errors on live production servers?",
         "Answer: We configured Dompdf with dynamic writable temporary directories (sys_get_temp_dir() and project cache fallback), proper chroot boundaries, and memory/time limit overrides. Additionally, we implemented a fail-safe exception handler: if PDF streaming fails on a restricted host, it seamlessly delivers a styled, self-contained printable HTML invoice with automatic window.print() execution."),

        ("Q4: What differentiates Lapify from generic e-commerce platforms like Amazon or OLX?",
         "Answer: Generic classifieds lack moderation and structured laptop specs, creating fraud risks. Lapify uniquely integrates specialized hardware filtering, mandatory admin quality inspection before listings go live, buyer protection upon delivery, and formal GST tax invoicing."),

        ("Q5: How is the database organized to handle multi-seller listings and order history?",
         "Answer: The relational database uses normalized tables including users, brands, brand_models, laptops, cart, orders, order_items, payments, and notifications with foreign key cascades and indexed lookup columns for high query performance.")
    ]

    for q_text, a_text in qa_list:
        add_body_p(q_text, bold_prefix="❓", italic=False)
        add_body_p(a_text, bold_prefix="💡", italic=False)
        doc.add_paragraph() # Spacer

    # ==========================================
    # SECTION 5: EXHIBITION CHECKLIST
    # ==========================================
    add_section_heading("📋 SECTION 5: EXHIBITION DAY CHECKLIST & TIPS")
    add_bullet("Pre-open 2 browser tabs: Tab 1 for Buyer storefront (`/buy.php`), Tab 2 for Admin Dashboard (`/admin/dashboard.php`).", bold_prefix="1. Tab Setup:")
    add_bullet("Ensure you have at least 1 laptop with stock = 10 and 1 listing in 'Pending Approval' status ready for live moderation demonstration.", bold_prefix="2. Data Readiness:")
    add_bullet("Speak slowly and clearly. Show genuine pride in your work — judges love passion, technical clarity, and well-structured code!", bold_prefix="3. Delivery:")

    doc.save(filename)
    print(f"Successfully created: {filename}")

if __name__ == "__main__":
    out_path = os.path.join(os.path.dirname(__file__), "..", "Lapify_Exhibition_Speech_and_Presentation_Guide.docx")
    out_path = os.path.abspath(out_path)
    create_speech_document(out_path)
